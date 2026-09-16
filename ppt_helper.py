import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 16:9 슬라이드 크기
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# 컬러 정의
C_NAVY_DARK = RGBColor(12, 27, 48)      # #0c1b30
C_BLUE_PRIMARY = RGBColor(0, 119, 200)  # #0077c8
C_BLUE_LIGHT = RGBColor(240, 247, 255)  # #f0f7ff
C_BLUE_BORDER = RGBColor(186, 230, 253) # #bae6fd
C_CARD_BG = RGBColor(248, 250, 252)     # #f8fafc
C_CARD_BORDER = RGBColor(226, 232, 240) # #e2e8f0
C_TEXT_DARK = RGBColor(15, 23, 42)      # #0f172a
C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748b
C_WHITE = RGBColor(255, 255, 255)
C_GREEN = RGBColor(5, 150, 105)         # #059669
C_AMBER = RGBColor(217, 119, 6)         # #d97706
C_RED = RGBColor(220, 38, 38)           # #dc2626
C_PURPLE = RGBColor(109, 40, 217)       # #6d28d9

def create_base_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs

def add_header_footer(slide, category, title, slide_num, total_slides):
    # 상단 헤더 바
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.name = 'Malgun Gothic'
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_BLUE_PRIMARY
    
    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.name = 'Malgun Gothic'
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = C_NAVY_DARK
    
    # 상단 구분선
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = C_BLUE_PRIMARY
    line.line.color.rgb = C_BLUE_PRIMARY
    
    # 하단 푸터 바
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
    tf_f = footer_box.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    p_f = tf_f.paragraphs[0]
    p_f.text = f"글로벌 세아(SAE-A) 감사 지적사항(CAP) 관리 시스템 사용자 매뉴얼   |   {slide_num} / {total_slides}"
    p_f.font.name = 'Malgun Gothic'
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = C_TEXT_MUTED

def add_title_slide(prs, main_title, sub_title, role_badge, role_desc):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # 배경 네이비 사각형
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY_DARK
    bg.line.fill.background()
    
    # 좌측 장식 라인
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.12), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_BLUE_PRIMARY
    bar.line.fill.background()
    
    # 텍스트 박스
    tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    # 뱃지
    p_badge = tf.paragraphs[0]
    p_badge.text = f"  {role_badge} 사용자 전용 매뉴얼  "
    p_badge.font.name = 'Malgun Gothic'
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = C_BLUE_PRIMARY
    
    # 메인 타이틀
    p_title = tf.add_paragraph()
    p_title.text = main_title
    p_title.font.name = 'Malgun Gothic'
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = C_WHITE
    p_title.space_before = Pt(14)
    p_title.space_after = Pt(10)
    
    # 서브 타이틀
    p_sub = tf.add_paragraph()
    p_sub.text = sub_title
    p_sub.font.name = 'Malgun Gothic'
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    
    # 대상 설명
    p_desc = tf.add_paragraph()
    p_desc.text = f"📌 적용 대상: {role_desc}"
    p_desc.font.name = 'Malgun Gothic'
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = RGBColor(148, 163, 184)
    p_desc.space_before = Pt(24)
    
    p_date = tf.add_paragraph()
    p_date.text = "글로벌 세아(SAE-A) 감사실 | Version 2.0"
    p_date.font.name = 'Malgun Gothic'
    p_date.font.size = Pt(11)
    p_date.font.color.rgb = RGBColor(100, 116, 139)
    p_date.space_before = Pt(10)

def add_cards_slide(prs, category, title, cards_info, slide_num, total_slides):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_header_footer(slide, category, title, slide_num, total_slides)
    
    num_cards = len(cards_info)
    gap = Inches(0.25)
    total_w = Inches(11.733)
    card_w = (total_w - gap * (num_cards - 1)) / num_cards
    card_h = Inches(5.2)
    top = Inches(1.6)
    
    for i, c in enumerate(cards_info):
        left = Inches(0.8) + i * (card_w + gap)
        
        # 카드 배경 상자
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = c.get('bg_color', C_CARD_BG)
        card.line.color.rgb = c.get('border_color', C_CARD_BORDER)
        card.line.width = Pt(1.5)
        
        # 카드 상단 헤더 바
        if 'header_color' in c:
            hb = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, Inches(0.55))
            hb.fill.solid()
            hb.fill.fore_color.rgb = c['header_color']
            hb.line.fill.background()
            
            tb_h = slide.shapes.add_textbox(left, top, card_w, Inches(0.55))
            tf_h = tb_h.text_frame
            tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_h = tf_h.paragraphs[0]
            p_h.text = c.get('header_title', '')
            p_h.font.name = 'Malgun Gothic'
            p_h.font.size = Pt(13)
            p_h.font.bold = True
            p_h.font.color.rgb = C_WHITE
            p_h.alignment = PP_ALIGN.CENTER
            
        # 카드 본문
        content_top = top + (Inches(0.65) if 'header_color' in c else Inches(0.2))
        content_h = card_h - (Inches(0.8) if 'header_color' in c else Inches(0.4))
        tb_c = slide.shapes.add_textbox(left + Inches(0.2), content_top, card_w - Inches(0.4), content_h)
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        
        if 'title' in c and 'header_color' not in c:
            p_ct = tf_c.paragraphs[0]
            p_ct.text = c['title']
            p_ct.font.name = 'Malgun Gothic'
            p_ct.font.size = Pt(15)
            p_ct.font.bold = True
            p_ct.font.color.rgb = c.get('title_color', C_NAVY_DARK)
            p_ct.space_after = Pt(10)
            
        for item in c.get('items', []):
            p = tf_c.add_paragraph() if ('title' in c and 'header_color' not in c) or tf_c.paragraphs[0].text else tf_c.paragraphs[0]
            if isinstance(item, tuple):
                strong_part, text_part = item
                r1 = p.add_run()
                r1.text = strong_part + " "
                r1.font.name = 'Malgun Gothic'
                r1.font.bold = True
                r1.font.size = Pt(11)
                r1.font.color.rgb = c.get('item_color', C_TEXT_DARK)
                
                r2 = p.add_run()
                r2.text = text_part
                r2.font.name = 'Malgun Gothic'
                r2.font.size = Pt(11)
                r2.font.color.rgb = C_TEXT_DARK
            else:
                r = p.add_run()
                r.text = item
                r.font.name = 'Malgun Gothic'
                r.font.size = Pt(11)
                r.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(8)

def add_table_slide(prs, category, title, headers, rows, col_widths, slide_num, total_slides):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    add_header_footer(slide, category, title, slide_num, total_slides)
    
    left = Inches(0.8)
    top = Inches(1.6)
    width = Inches(11.733)
    height = Inches(0.6 * (len(rows) + 1))
    
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    
    # 열 너비 설정
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
        
    # 헤더 행
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = 'Malgun Gothic'
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER
        
    # 데이터 행
    for r_idx, row in enumerate(rows):
        bg_c = RGBColor(255, 255, 255) if r_idx % 2 == 0 else RGBColor(248, 250, 252)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_c
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.name = 'Malgun Gothic'
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT_DARK
            if c_idx == 0:
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER

print("Base helper loaded successfully")
